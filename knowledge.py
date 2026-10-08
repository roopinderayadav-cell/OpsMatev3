"""
knowledge.py - OpsMate's knowledge base and answer engine.

  * stores documents per region (SQLite + files on disk), reads PDF/Word/PPT/Excel,
    tables, and (with AI) screenshots/images/scans inside documents
  * Layer 2 search: word matching + Gemini semantic search across ALL documents
  * runs the 3 layers for every question and decides: answered / not found / conflict
  * SELF-LEARNING:
      - 👍 answers are saved automatically and reused for similar questions
      - 👍 / 👎 on a passage raises / lowers it for similar questions in future
      - 👎 on an auto-learned answer switches that answer off
      - after a new upload, open knowledge gaps are re-checked automatically
  * knowledge gaps, analytics data and the Excel export
"""
import io
import json
import re
import sqlite3
import time
import uuid
from datetime import datetime, timedelta

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

import brain
import config
import llm

# ===========================================================================
# 1. DATABASE
# ===========================================================================

def _conn():
    config.DATA_DIR.mkdir(exist_ok=True)
    config.UPLOAD_DIR.mkdir(exist_ok=True)
    con = sqlite3.connect(config.DB_PATH, check_same_thread=False, timeout=30)
    con.row_factory = sqlite3.Row
    return con


def _add_cols(con, table, cols):
    have = [r[1] for r in con.execute(f"PRAGMA table_info({table})")]
    for name, kind in cols:
        if name not in have:
            con.execute(f"ALTER TABLE {table} ADD COLUMN {name} {kind}")


def init_db():
    with _conn() as con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, doc_type TEXT, title TEXT, filename TEXT,
                stored_path TEXT, uploaded_by TEXT, uploaded_at TEXT,
                version TEXT, chunks INTEGER DEFAULT 0
            );
            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                doc_id INTEGER, region TEXT, location TEXT, text TEXT
            );
            CREATE TABLE IF NOT EXISTS questions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, asked_by TEXT, question TEXT, answer TEXT,
                sources TEXT, score REAL, feedback INTEGER, comment TEXT,
                status TEXT, asked_at TEXT
            );
            CREATE TABLE IF NOT EXISTS learned (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, question TEXT, answer TEXT,
                added_by TEXT, added_at TEXT, from_question INTEGER
            );
            CREATE TABLE IF NOT EXISTS feedback_pairs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                region TEXT, question TEXT, source_key TEXT, value INTEGER, at TEXT
            );
            CREATE TABLE IF NOT EXISTS announcements (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT, body TEXT, level TEXT, region TEXT,
                author TEXT, created_at TEXT, active INTEGER DEFAULT 1
            );
            """
        )
        # upgrades from Compass v1 / v2 - existing data is kept
        _add_cols(con, "documents", [("status", "TEXT DEFAULT 'Active'"), ("owner", "TEXT"),
                                     ("keywords", "TEXT")])
        _add_cols(con, "chunks", [("embedding", "BLOB")])
        _add_cols(con, "questions", [("image_path", "TEXT"), ("meta", "TEXT"), ("topic", "TEXT"),
                                     ("result", "TEXT"), ("confidence", "TEXT"),
                                     ("gap_status", "TEXT"), ("gap_owner", "TEXT"),
                                     ("gap_note", "TEXT")])
        _add_cols(con, "learned", [("source", "TEXT DEFAULT 'verified'"),
                                   ("active", "INTEGER DEFAULT 1"), ("votes", "INTEGER DEFAULT 0"),
                                   ("embedding", "BLOB")])


def _now():
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def _vec_to_blob(v):
    return np.asarray(v, dtype=np.float32).tobytes()


def _blob_to_vec(b):
    return np.frombuffer(b, dtype=np.float32) if b else None


# ===========================================================================
# 2. TEXT EXTRACTION
# ===========================================================================

IMAGE_EXTS = ("png", "jpg", "jpeg")


def extract_text(filename: str, data: bytes):
    """Return (parts, images).
    parts  = list of (location, text), e.g. ("2. Create a New PNR", "...")
    images = list of (location, image_bytes) - screenshots/tables inside the file."""
    ext = filename.rsplit(".", 1)[-1].lower()
    parts, images = [], []

    if ext == "pdf":
        try:
            import pdfplumber
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for i, page in enumerate(pdf.pages, 1):
                    txt = page.extract_text() or ""
                    parts.append((f"Page {i}", txt))
                    for t_i, table in enumerate(page.extract_tables() or [], 1):
                        rows = [" | ".join((c or "").strip().replace("\n", " ") for c in r) for r in table if r]
                        if len(rows) > 1:
                            parts.append((f"Page {i} - Table {t_i}", "\n".join(rows)))
                    if len(txt.strip()) < 30:          # scanned page: read the picture
                        images.append((f"Page {i} (scanned)", _render_pdf_page(data, i - 1)))
        except ImportError:
            from pypdf import PdfReader
            for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
                parts.append((f"Page {i}", page.extract_text() or ""))
        try:                                            # pictures placed in the PDF
            from pypdf import PdfReader
            for i, page in enumerate(PdfReader(io.BytesIO(data)).pages, 1):
                for img in page.images:
                    images.append((f"Page {i}", img.data))
        except Exception:
            pass

    elif ext == "docx":
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = docx.Document(io.BytesIO(data))
        list_kind = _docx_list_reader(d)
        section, buf, t_i, counter = "Document", [], 0, {}
        for el in d.element.body.iterchildren():
            tag = el.tag.split("}")[-1]
            if tag == "p":
                p = Paragraph(el, d)
                style = (p.style.name if p.style is not None else "").lower()
                if (style.startswith("heading") or style == "title") and p.text.strip():
                    if buf:
                        parts.append((section, "\n".join(buf)))
                    section, buf, counter = p.text.strip(), [p.text.strip()], {}
                elif p.text.strip():
                    kind, key = list_kind(p)
                    text = p.text.strip()
                    if kind == "decimal":       # Word's automatic 1. 2. 3. is not in the text
                        counter[key] = counter.get(key, 0) + 1
                        if not re.match(r"^\d+[.)]", text):
                            text = f"{counter[key]}. {text}"
                    elif kind == "bullet" and not text.startswith(("-", "•")):
                        text = f"- {text}"
                    buf.append(text)
                for blip in el.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip"):
                    rid = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                    if rid and rid in d.part.related_parts:
                        images.append((section, d.part.related_parts[rid].blob))
            elif tag == "tbl":
                t_i += 1
                table = Table(el, d)
                rows = []
                for r in table.rows:
                    cells = []
                    for c in r.cells:
                        if not cells or c.text.strip() != cells[-1]:   # merged cells repeat
                            cells.append(c.text.strip())
                    rows.append(" | ".join(cells))
                parts.append((f"{section} - Table {t_i}", "\n".join(rows)))
        if buf:
            parts.append((section, "\n".join(buf)))

    elif ext == "pptx":
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs = Presentation(io.BytesIO(data))
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for sh in slide.shapes:
                if sh.has_text_frame:
                    texts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    texts += [" | ".join(c.text for c in r.cells) for r in sh.table.rows]
                if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    images.append((f"Slide {i}", sh.image.blob))
            parts.append((f"Slide {i}", "\n".join(texts)))

    elif ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    rows.append(" | ".join(cells))
            parts.append((f"Sheet {ws.title}", "\n".join(rows)))

    elif ext in IMAGE_EXTS:                             # a screenshot / table uploaded on its own
        images.append(("Image", data))

    else:  # txt, md, csv - split into sections at heading lines
        parts += _text_sections(data.decode("utf-8", errors="ignore"))

    return [(loc, txt) for loc, txt in parts if txt and txt.strip()], images


def _docx_list_reader(d):
    """Returns a function telling whether a Word paragraph is a numbered or bulleted list item."""
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    fmts = {}
    try:
        numbering = d.part.numbering_part.element
        abstract = {a.get(W + "abstractNumId"): a for a in numbering.iter(W + "abstractNum")}
        for num in numbering.iter(W + "num"):
            ref = num.find(W + "abstractNumId")
            a = abstract.get(ref.get(W + "val")) if ref is not None else None
            if a is None:
                continue
            for lvl in a.iter(W + "lvl"):
                f = lvl.find(W + "numFmt")
                fmts[(num.get(W + "numId"), lvl.get(W + "ilvl"))] = f.get(W + "val") if f is not None else ""
    except Exception:
        pass

    def kind(p):
        ppr = p._p.pPr
        num = ppr.numPr if ppr is not None else None
        style = (p.style.name if p.style is not None else "").lower()
        if num is None or num.numId is None:
            if "list number" in style:
                return "decimal", style
            if "list bullet" in style:
                return "bullet", style
            return None, None
        num_id = str(num.numId.val)
        ilvl = str(num.ilvl.val) if num.ilvl is not None else "0"
        fmt = fmts.get((num_id, ilvl), "decimal")
        return ("bullet" if fmt == "bullet" else "decimal"), (num_id, ilvl)
    return kind


def _is_heading(line, nxt):
    t = line.strip()
    if not t or not nxt.strip():
        return False
    if t.startswith("#"):
        return True
    if t.startswith(("Q:", "A:", "-", "*", "•")) or re.match(r"^\d+[.)]\s", t) or "|" in t:
        return False
    return len(t.split()) <= 7 and not t.endswith((".", ":", ";", ",", "?")) and t[0].isupper()


def _text_sections(text):
    """Plain text / markdown: short title lines (or # headings) start a new section."""
    lines = text.splitlines()
    out, section, buf = [], "Text", []
    for i, line in enumerate(lines):
        nxt = lines[i + 1] if i + 1 < len(lines) else ""
        if _is_heading(line, nxt):
            if buf:
                out.append((section, "\n".join(buf)))
            section = line.strip().lstrip("#").strip()
            buf = [section]
        elif line.strip():
            buf.append(line.strip())
    if buf:
        out.append((section, "\n".join(buf)))
    return out


def _render_pdf_page(data, index):
    try:
        import pypdfium2 as pdfium
        page = pdfium.PdfDocument(data)[index]
        buf = io.BytesIO()
        page.render(scale=2).to_pil().save(buf, format="PNG")
        return buf.getvalue()
    except Exception:
        return b""


def read_images(title, images):
    """Turn images inside a document into searchable text (needs the AI key).
    Runs several images at the same time so uploads stay quick."""
    if not images or not llm.ai_on() or not config.READ_DOC_IMAGES:
        return [], len(images)
    from concurrent.futures import ThreadPoolExecutor
    images = [(loc, b) for loc, b in images if b][: config.MAX_IMAGES_PER_DOC]

    def work(item):
        loc, b = item
        try:
            return loc, brain.describe_doc_image(b, title, loc)
        except Exception:
            return loc, ""

    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(work, images))
    out, counter = [], {}
    for loc, text in results:
        if text:
            counter[loc] = counter.get(loc, 0) + 1
            out.append((f"{loc} - image {counter[loc]}", "[From image in document]\n" + text))
    return out, 0


def chunk_text(text: str, size=config.CHUNK_SIZE, overlap=config.CHUNK_OVERLAP):
    """Split long text into overlapping pieces that keep paragraphs together.
    FAQ-style text ("Q: ..." lines) is split so each question+answer is its own piece."""
    text = re.sub(r"[ \t]+", " ", text).strip()
    faq = re.split(r"\n(?=\s*(?:Q\s*[:.\-]|Question\s*\d*\s*[:.]))", text, flags=re.I)
    if len(faq) >= 3:
        out = []
        for piece in faq:
            out += chunk_text(piece, size, overlap) if len(piece) > size else [piece.strip()]
        return [o for o in out if o]
    if len(text) <= size:
        return [text]
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n|\n", text) if p.strip()]
    chunks, current = [], ""
    for p in paragraphs:
        if len(current) + len(p) + 1 <= size:
            current = f"{current}\n{p}".strip()
        else:
            if current:
                chunks.append(current)
            while len(p) > size:          # very long paragraph
                chunks.append(p[:size])
                p = p[size - overlap:]
            current = (current[-overlap:] + "\n" + p).strip() if current else p
    if current:
        chunks.append(current)
    return chunks


# ===========================================================================
# 3. DOCUMENT MANAGEMENT
# ===========================================================================

def add_document(region, doc_type, title, filename, data: bytes, uploaded_by, version="", owner=""):
    """Returns (doc_id, pieces, info). info = images read/skipped, semantic index, gaps re-checked."""
    parts, images = extract_text(filename, data)
    image_parts, skipped = read_images(title or filename, images)
    parts += image_parts
    if not parts:
        if images and not llm.ai_on():
            raise ValueError("This file is only images/scans. Reading images needs AI mode "
                             "(add GEMINI_API_KEY).")
        raise ValueError("No readable text found.")

    stored = config.UPLOAD_DIR / f"{uuid.uuid4().hex[:8]}_{filename}"
    stored.write_bytes(data)

    # template guidance sections are never used for answers
    parts = [(loc, "\n".join(l for l in txt.splitlines() if not l.strip().lower().startswith("replace the example")))
             for loc, txt in parts if not loc.lower().startswith(("how to use this template", "instructions"))]
    # the "Document information" table becomes hidden keywords for the whole document
    keywords = " ".join(txt.replace("|", " ") for loc, txt in parts
                        if loc.lower().startswith("document information"))
    parts = [(loc, txt) for loc, txt in parts if not loc.lower().startswith("document information")]
    parts = [(loc, txt) for loc, txt in parts if txt.strip()]
    pieces = []
    for loc, txt in parts:
        for piece in ([txt] if "Table" in loc or "image" in loc else chunk_text(txt)):
            if len(piece) < 40 and len(parts) > 1:      # skip title-only fragments
                continue
            pieces.append((loc, piece))

    vectors, semantic = None, "off"
    if llm.embeddings_on() and pieces:
        try:
            vectors = llm.embed([f"{title or filename} | {loc}\n{p}" for loc, p in pieces])
            semantic = "on"
        except llm.AIError as e:
            semantic = f"failed: {e}"

    with _conn() as con:
        cur = con.execute(
            "INSERT INTO documents (region, doc_type, title, filename, stored_path, uploaded_by,"
            " uploaded_at, version, status, owner, keywords) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (region, doc_type, title or filename, filename, str(stored), uploaded_by, _now(),
             version, "Active", owner, keywords[:2000]))
        doc_id = cur.lastrowid
        for i, (loc, piece) in enumerate(pieces):
            con.execute("INSERT INTO chunks (doc_id, region, location, text, embedding) VALUES (?,?,?,?,?)",
                        (doc_id, region, loc, piece,
                         _vec_to_blob(vectors[i]) if vectors is not None else None))
        con.execute("UPDATE documents SET chunks=? WHERE id=?", (len(pieces), doc_id))
    rechecked = recheck_gaps(region, doc_title=title or filename)
    return doc_id, len(pieces), {"images_found": len(images), "images_read": len(image_parts),
                                 "images_skipped": skipped, "semantic": semantic,
                                 "gaps_rechecked": rechecked}


def reindex_embeddings(progress=None):
    """Build the semantic index for pieces/answers that don't have one yet (after adding a key)."""
    if not llm.embeddings_on():
        raise llm.AIError("Add GEMINI_API_KEY first")
    with _conn() as con:
        rows = con.execute("SELECT c.id, c.location, c.text, d.title FROM chunks c JOIN documents d "
                           "ON d.id=c.doc_id WHERE c.embedding IS NULL").fetchall()
        learned = con.execute("SELECT id, question, answer FROM learned WHERE embedding IS NULL").fetchall()
    done = 0
    for i in range(0, len(rows), 50):
        batch = rows[i:i + 50]
        vecs = llm.embed([f"{r['title']} | {r['location']}\n{r['text']}" for r in batch])
        with _conn() as con:
            for r, v in zip(batch, vecs):
                con.execute("UPDATE chunks SET embedding=? WHERE id=?", (_vec_to_blob(v), r["id"]))
        done += len(batch)
        if progress:
            progress(done / max(1, len(rows)))
    if learned:
        vecs = llm.embed([f"{l['question']}\n{l['answer']}" for l in learned])
        with _conn() as con:
            for l, v in zip(learned, vecs):
                con.execute("UPDATE learned SET embedding=? WHERE id=?", (_vec_to_blob(v), l["id"]))
    return done + len(learned)


def semantic_coverage():
    with _conn() as con:
        total = con.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
        with_vec = con.execute("SELECT COUNT(*) FROM chunks WHERE embedding IS NOT NULL").fetchone()[0]
    return with_vec, total


def load_samples_if_empty():
    """Load sample_docs/ when the knowledge base is empty. Region = start of the file name
    (NORAM_, EUROPE_, APAC_, GLOBAL_); type = FAQ / Policy / Quick Guide / SOP by name."""
    with _conn() as con:
        if con.execute("SELECT COUNT(*) FROM documents").fetchone()[0]:
            return
    folder = config.BASE_DIR / "sample_docs"
    if not folder.exists():
        return
    regions = config.REGIONS + [config.GLOBAL_REGION]
    for f in sorted(folder.iterdir()):
        region = f.name.split("_", 1)[0].upper()
        if region not in regions or f.suffix.lstrip(".").lower() not in config.ALLOWED_EXTENSIONS:
            continue
        name = f.name.lower()
        doc_type = ("FAQ" if "faq" in name else "Policy" if "policy" in name
                    else "Quick Guide" if "guide" in name else "SOP")
        title = f.stem.split("_", 1)[-1].replace("_", " ")
        try:
            add_document(region, doc_type, title, f.name, f.read_bytes(), "Sample data", "1.0")
        except Exception:
            pass


def delete_document(doc_id):
    with _conn() as con:
        row = con.execute("SELECT stored_path FROM documents WHERE id=?", (doc_id,)).fetchone()
        con.execute("DELETE FROM chunks WHERE doc_id=?", (doc_id,))
        con.execute("DELETE FROM documents WHERE id=?", (doc_id,))
    if row:
        try:
            from pathlib import Path
            Path(row["stored_path"]).unlink(missing_ok=True)
        except OSError:
            pass


def set_document_status(doc_id, status):
    """Active documents are searched; Archived ones are kept but ignored."""
    with _conn() as con:
        con.execute("UPDATE documents SET status=? WHERE id=?", (status, doc_id))


def list_documents(region=None, include_global=True, status=None):
    q, args, where = "SELECT * FROM documents", [], []
    if region:
        regions = [region, config.GLOBAL_REGION] if include_global else [region]
        where.append(f"region IN ({','.join('?' * len(regions))})")
        args += regions
    if status:
        where.append("COALESCE(status,'Active')=?")
        args.append(status)
    if where:
        q += " WHERE " + " AND ".join(where)
    q += " ORDER BY region, doc_type, title"
    with _conn() as con:
        return [dict(r) for r in con.execute(q, args)]


def get_document(doc_id):
    with _conn() as con:
        r = con.execute("SELECT * FROM documents WHERE id=?", (doc_id,)).fetchone()
    return dict(r) if r else None


def document_text(doc_id, limit=6000):
    with _conn() as con:
        rows = con.execute("SELECT location, text FROM chunks WHERE doc_id=? ORDER BY id", (doc_id,)).fetchall()
    return "\n\n".join(r["text"] for r in rows)[:limit]


def popular_questions(region, limit=6):
    """Most-asked questions that OpsMate answered well - used for 'Try asking'."""
    with _conn() as con:
        rows = con.execute(
            "SELECT question, COUNT(*) n FROM questions WHERE region=? AND result='answered' "
            "AND COALESCE(feedback,0)>=0 AND question NOT LIKE '%[screenshot]' "
            "GROUP BY lower(question) ORDER BY n DESC, MAX(id) DESC LIMIT ?", (region, limit)).fetchall()
    return [r["question"] for r in rows]


def read_file(doc):
    from pathlib import Path
    p = Path(doc["stored_path"])
    return p.read_bytes() if p.exists() else b""


# ===========================================================================
# 4. SEARCH  (Layer 2)
# ===========================================================================

def _candidates(region):
    """Everything searchable for a region: active docs + GLOBAL docs + learned answers."""
    regions = (region, config.GLOBAL_REGION)
    with _conn() as con:
        rows = con.execute(
            "SELECT c.id, c.text, c.location, c.embedding, d.id AS doc_id, d.title, d.doc_type, "
            "d.region, d.version, d.keywords FROM chunks c JOIN documents d ON d.id = c.doc_id "
            "WHERE c.region IN (?,?) AND COALESCE(d.status,'Active')='Active'", regions).fetchall()
        learned = con.execute(
            "SELECT id, question, answer, added_by, added_at, region, source, embedding FROM learned "
            "WHERE region IN (?,?) AND COALESCE(active,1)=1", regions).fetchall()
    items = [
        {"kind": "doc", "text": r["text"], "title": r["title"], "type": r["doc_type"],
         "location": r["location"], "region": r["region"], "version": r["version"],
         "doc": f"d{r['doc_id']}", "doc_id": r["doc_id"], "chunk_id": r["id"],
         "key": f"c{r['id']}", "vec": _blob_to_vec(r["embedding"]),
         # the section heading counts twice: headings are a strong signal
         "search": f"{r['title']} {r['location']} {r['location']} {r['text']}",
         "keywords": r["keywords"] or ""}
        for r in rows]
    for l in learned:
        auto = (l["source"] or "verified") == "auto"
        items.append(
            {"kind": "learned", "auto": auto, "text": l["answer"],
             "title": "Learned from staff feedback" if auto else "Verified answer",
             "type": "Auto-learned Q&A" if auto else "Verified Q&A",
             "location": f"{'learned' if auto else 'added by ' + (l['added_by'] or '')} on {l['added_at']}",
             "region": l["region"], "question": l["question"], "doc": f"l{l['id']}",
             "key": f"l{l['id']}", "vec": _blob_to_vec(l["embedding"]),
             "search": f"{l['question']} {l['question']} {l['answer']}"})
    return items


# Small stop-word list on purpose: words like "name", "change", "first" matter in SOPs.
STOP_WORDS = """a an the is are was were be been am do does did i me my we our you your it its
of to in on for at by with from and or as this that these those can could should would will
how what when where which who why please tell know need want get about any there""".split()

_PLURAL = re.compile(r"\b([a-z]{3,}[^s\W])s\b")


def _norm(text):
    """lower-case, '2,000' -> '2000', simple singular form ('commands' -> 'command')."""
    return _PLURAL.sub(r"\1", re.sub(r"(\d),(\d{3})", r"\1\2", text.lower()))


def _feedback_multipliers(region, query):
    """SELF-LEARNING: passages staff marked 👍 / 👎 for SIMILAR questions are
    moved up (x1.3 per vote, max x1.9) or down (x0.5 per vote)."""
    with _conn() as con:
        rows = con.execute("SELECT question, source_key, value FROM feedback_pairs "
                           "WHERE region IN (?,?)", (region, config.GLOBAL_REGION)).fetchall()
    if not rows:
        return {}
    texts = [r["question"] for r in rows]
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(texts + [query])
    sims = cosine_similarity(vec.transform([query]), vec.transform(texts)).ravel()
    mult = {}
    for r, s in zip(rows, sims):
        if s >= config.FEEDBACK_SIMILARITY:
            m = mult.get(r["source_key"], 1.0)
            mult[r["source_key"]] = min(1.9, m * 1.3) if r["value"] > 0 else m * 0.5
    return mult


def search(region, queries, k=config.CANDIDATES, fmt="short", query_vec=None):
    """Search ALL documents with every phrasing of the question.
    score = word match (and, with Gemini, meaning match), adjusted by staff feedback.
    Keeps at most PER_DOC passages from any one document so answers can be compared."""
    if isinstance(queries, str):
        queries = [queries]
    queries = [q for q in dict.fromkeys(queries) if q and q.strip()]
    items = _candidates(region)
    if not items or not queries:
        return []
    texts = [i["search"] + " " + i.get("keywords", "") for i in items] + queries
    words = TfidfVectorizer(stop_words=STOP_WORDS, ngram_range=(1, 2), sublinear_tf=True,
                            preprocessor=_norm)
    pieces = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), sublinear_tf=True)
    mw, mp = words.fit_transform(texts), pieces.fit_transform(texts)
    n = len(items)
    heads = words.transform([f"{i['title']} {i.get('location', '')} {i.get('question', '')} "
                             + (i["text"].split("\n", 1)[0] if i["text"].lstrip().startswith("Q:") else "")
                             for i in items])
    lex = None
    for qi in range(len(queries)):
        s = (0.6 * cosine_similarity(mw[n + qi], mw[:n]).ravel()
             + 0.4 * cosine_similarity(mp[n + qi], mp[:n]).ravel()
             + 0.3 * cosine_similarity(mw[n + qi], heads).ravel())
        lex = s if lex is None else np.maximum(lex, s)

    sem = None
    if query_vec is not None and any(i["vec"] is not None for i in items):
        sem = np.array([float(np.dot(query_vec, i["vec"])) if i["vec"] is not None
                        and len(i["vec"]) == len(query_vec) else 0.0 for i in items])

    mult = _feedback_multipliers(region, queries[0])
    # SELF-LEARNING: a learned answer whose question is worded like this one wins clearly
    learned_idx = [i for i, it in enumerate(items) if it["kind"] == "learned"]
    if learned_idx:
        qv = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(
            [items[i]["question"] for i in learned_idx] + queries)
        lq = qv.transform([items[i]["question"] for i in learned_idx])
        qsim = cosine_similarity(qv.transform(queries), lq).max(axis=0)
        for i, sim in zip(learned_idx, qsim):
            items[i]["qsim"] = float(sim)
    for idx, item in enumerate(items):
        boost = 1.0
        if item["kind"] == "learned":
            boost = config.AUTO_LEARNED_BOOST if item.get("auto") else config.LEARNED_BOOST
            if item.get("qsim", 0) >= 0.6:
                boost *= 1.6
        if fmt == "table" and "Table" in item.get("location", ""):
            boost *= 1.3                     # asked for a list/table -> prefer tables
        if fmt == "steps" and len(re.findall(r"(?:^|\n)\s*\d+[.)]\s", item["text"])) >= 3:
            boost *= 1.5                     # asked for steps -> prefer real numbered procedures
        elif fmt == "steps" and re.search(r"(^|\s)[2-9][.)]\s", item["text"]):
            boost *= 1.2
        boost *= mult.get(item["key"], 1.0)
        item["lex"] = float(lex[idx])
        item["sem"] = float(sem[idx]) if sem is not None else None
        base = item["lex"] if sem is None else 0.7 * item["sem"] + 0.3 * min(1.0, item["lex"] * 2.5)
        item["score"] = base * boost
        item["feedback"] = mult.get(item["key"], 1.0)
    ranked = sorted(items, key=lambda x: x["score"], reverse=True)
    top = ranked[0]["score"] if ranked else 0
    picked, per_doc = [], {}
    for r in ranked:
        if r["score"] <= 0 or r["score"] < top * 0.35 or len(picked) >= k:
            break
        if per_doc.get(r["doc"], 0) >= config.PER_DOC:
            continue
        per_doc[r["doc"]] = per_doc.get(r["doc"], 0) + 1
        picked.append(r)
    return picked


def expand_sections(hits, limit=3000):
    """A long section can be split into several pieces. Re-join the pieces that
    directly continue the matched one, so procedures come out complete (all steps).
    FAQ entries (Q: ...) are never merged with their neighbours."""
    out, seen = [], set()
    with _conn() as con:
        for h in hits:
            if h["kind"] != "doc":
                out.append(h); continue
            if h["chunk_id"] in seen:
                continue
            seen.add(h["chunk_id"])
            if h["text"].lstrip().startswith("Q:"):
                out.append(h); continue
            rows = con.execute("SELECT id, text FROM chunks WHERE doc_id=? AND location=? ORDER BY id",
                               (h["doc_id"], h["location"])).fetchall()
            ids = [r["id"] for r in rows]
            if len(rows) < 2 or h["chunk_id"] not in ids:
                out.append(h); continue
            i = j = ids.index(h["chunk_id"])
            while i > 0 and ids[i - 1] == ids[i] - 1 and not rows[i]["text"].lstrip().startswith("Q:"):
                i -= 1
            while j < len(ids) - 1 and ids[j + 1] == ids[j] + 1 and not rows[j + 1]["text"].lstrip().startswith("Q:"):
                j += 1
            lines = []
            for r in rows[i:j + 1]:
                seen.add(r["id"])
                for line in r["text"].splitlines():
                    line = line.strip()
                    if line and not any(line in l or l in line for l in lines[-6:] if len(l) > 20):
                        lines.append(line)
            out.append(dict(h, text="\n".join(lines)[:limit]))
    return out


def vocabulary(region):
    return brain.build_vocab(i["search"] for i in _candidates(region))


# ===========================================================================
# 5. ANSWERING  (the three layers)
# ===========================================================================

def _public(h):
    return {"title": h["title"], "type": h["type"], "location": h["location"],
            "region": h["region"], "score": round(h["score"], 3), "text": h["text"],
            "key": h["key"], "kind": h["kind"], "doc_id": h.get("doc_id")}


def ask(region, question, asked_by, image=None):
    """question: what the staff member typed (can be empty if they only pasted a screenshot)
    image: bytes of a pasted screenshot, or None. Returns a dict for the screen."""
    t0 = time.time()
    question = (question or "").strip()
    meta = {"layers": [], "notes": []}
    use_ai = llm.ai_on()

    # --- Screenshot: read it first so its text can be searched ---------------
    shot_text, shot_issue = "", ""
    if image:
        if use_ai:
            try:
                shot = brain.read_screenshot(image, question)
                shot_text, shot_issue = shot["text"], shot["issue"]
                meta["layers"].append("read screenshot")
            except llm.AIError as e:
                meta["notes"].append(str(e))
        else:
            meta["notes"].append("Screenshots are read only in AI mode (add GEMINI_API_KEY).")
    if not question:
        question = shot_issue or "What does this screenshot mean and what should I do?"

    # --- Layer 1: understand ---------------------------------------------------
    fixed, changes = brain.fix_spelling(question, vocabulary(region))
    fmt = brain.detect_format(fixed)
    topic = brain.detect_topic(fixed + " " + shot_text[:200])
    queries = [brain.search_text(question), brain.search_text(fixed)] + ([shot_text[:400]] if shot_text else [])
    meta["layers"].append("spelling + topic")

    qvec = None
    if llm.embeddings_on():
        try:
            qvec = llm.embed([fixed + (" " + shot_issue if shot_issue else "")], query=True)[0]
            meta["layers"].append("semantic search")
        except llm.AIError as e:
            meta["notes"].append(str(e))

    hits = search(region, queries, fmt=fmt, query_vec=qvec)
    best = hits[0] if hits else None
    weak = (not best) or (best["lex"] < config.CONF_MEDIUM and (best["sem"] or 0) < config.SEM_MEDIUM)
    if use_ai and (weak or image) and qvec is None:
        try:                                # only when the question is unclear - saves quota
            u = brain.ai_understand(question, shot_text)
            queries += [brain.search_text(u["corrected"])] + u["queries"]
            fmt = u["format"] if fmt == "short" else fmt
            if u["corrected"].lower() != question.lower():
                fixed, changes = u["corrected"], changes or [("", "")]
            hits = search(region, queries, fmt=fmt, query_vec=qvec)
            meta["layers"].append("AI rephrase")
        except llm.AIError as e:
            meta["notes"].append(str(e))

    # --- Layers 2 + 3: compare and answer ---------------------------------------
    hits = expand_sections(hits)
    st = brain.core_text(fixed)
    # re-rank the leaders: a passage that covers more of the question's key words wins
    for h in hits[:5]:
        h["cover"] = brain.coverage(st, h.get("search", h["text"]))
        h["score"] = h["score"] * (0.6 + 0.4 * h["cover"])
    hits = sorted(hits[:5], key=lambda x: x["score"], reverse=True) + hits[5:]
    best = hits[0] if hits else None
    cover = brain.coverage(st, best.get("search", best["text"])) if best else 0
    conf = brain.offline_confidence(best["lex"], cover, best["sem"], len(brain.content_words(st))) if best else "low"
    if best and best["kind"] == "learned":   # short, curated answers: trust a strong match
        if conf == "medium" or (conf == "low" and best["lex"] >= 0.25 and cover >= 0.5):
            conf = "high" if conf == "medium" else "medium"
    result, answer, important, conflict_pair, conflict_text, mode = "not_found", None, "", None, "", "documents"

    if hits and use_ai and (conf != "low" or (best["lex"] >= config.CONF_MEDIUM * 0.6)):
        try:
            a = brain.ai_answer(fixed if changes else question, region, hits, fmt, shot_text)
            mode = "ai"
            meta["layers"].append("AI compare + answer")
            result, conf = a["status"], a["confidence"]
            answer, important, conflict_text = a["answer"], a["important"], a["conflict"]
            if result == "answered" and conf == "low":
                result = "not_found"
            if a["used"]:                   # show the sources the AI actually used first
                used = [hits[i - 1] for i in a["used"] if 0 < i <= len(hits)]
                hits = used + [h for h in hits if h not in used]
            if result == "conflict":
                conflict_pair = [h for h in hits if h["kind"] == "doc"][:2]
        except llm.AIError as e:
            meta["notes"].append(str(e))
            mode = "documents"

    if mode == "documents" and hits:
        # without an AI answer, use the tested word-match rule to decide (not the semantic one)
        conf = brain.offline_confidence(best["lex"], cover, None, len(brain.content_words(st)))
        if best["kind"] == "learned" and (conf == "medium" or (best["lex"] >= 0.25 and cover >= 0.5)):
            conf = "high" if conf == "medium" else "medium"
        pair = brain.offline_conflict(st, hits) if conf != "low" else None
        if pair and best["kind"] != "learned":
            result, conflict_pair = "conflict", list(pair)
        elif conf != "low":
            result, answer = "answered", brain.crisp_answer(fixed, hits, fmt)

    if result == "not_found":
        answer = config.FALLBACK_TEXT
    elif result == "conflict":
        answer = config.CONFLICT_TEXT

    sources = [_public(h) for h in (conflict_pair or hits)] if result != "not_found" else []
    meta.update({"mode": mode, "format": fmt, "topic": topic, "seconds": round(time.time() - t0, 1),
                 "did_you_mean": fixed if (changes and fixed.lower() != question.lower()) else "",
                 "documents": len({s["title"] for s in sources}), "important": important,
                 "conflict_text": conflict_text, "screenshot_text": shot_text[:600],
                 "learned": bool(best and best["kind"] == "learned" and result == "answered")})
    status = "answered" if result == "answered" else "needs_review"
    with _conn() as con:
        cur = con.execute(
            "INSERT INTO questions (region, asked_by, question, answer, sources, score, status, asked_at,"
            " meta, topic, result, confidence, gap_status) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (region, asked_by, question + (" [screenshot]" if image else ""), answer,
             json.dumps(sources), best["score"] if best else 0.0, status, _now(), json.dumps(meta),
             topic, result, conf, None if result == "answered" else "Open"))
        qid = cur.lastrowid
        if image:
            path = config.UPLOAD_DIR / f"question_{qid}.png"
            path.write_bytes(image)
            con.execute("UPDATE questions SET image_path=? WHERE id=?", (str(path), qid))
    return {"id": qid, "answer": answer, "sources": sources, "result": result,
            "found": result == "answered", "confidence": conf, "meta": meta, "question": question}


# ===========================================================================
# 6. LEARNING LOOP
# ===========================================================================

def give_feedback(qid, helpful: bool, comment="", trusted=False):
    """👍 / 👎 from staff. This is where OpsMate learns by itself."""
    with _conn() as con:
        q = con.execute("SELECT * FROM questions WHERE id=?", (qid,)).fetchone()
        if not q:
            return
        con.execute("UPDATE questions SET feedback=?, comment=? WHERE id=?",
                    (1 if helpful else -1, comment, qid))
        if not helpful:
            con.execute("UPDATE questions SET status='needs_review', gap_status=COALESCE(gap_status,'Open') "
                        "WHERE id=?", (qid,))
        sources = json.loads(q["sources"] or "[]")
        question = re.sub(r"\s*\[screenshot\]$", "", q["question"])
        for s in sources[:2]:              # the passages that produced the answer
            if s.get("key"):
                con.execute("INSERT INTO feedback_pairs (region, question, source_key, value, at) "
                            "VALUES (?,?,?,?,?)", (q["region"], question, s["key"],
                                                   1 if helpful else -1, _now()))
        if not helpful:                    # a wrong auto-learned answer is switched off
            for s in sources[:1]:
                if s.get("key", "").startswith("l") and s.get("type") == "Auto-learned Q&A":
                    con.execute("UPDATE learned SET active=0 WHERE id=? AND source='auto'",
                                (int(s["key"][1:]),))
    if helpful and config.AUTO_LEARN and q["result"] == "answered":
        _auto_learn(q, question, trusted)


def _auto_learn(q, question, trusted=False):
    """Save an answer staff found helpful so similar questions get it straight away."""
    if not sources_from_documents(q):
        return                             # already a learned answer - nothing new to save
    with _conn() as con:
        voters = con.execute("SELECT COUNT(DISTINCT asked_by) FROM questions WHERE region=? AND feedback=1 "
                             "AND lower(question)=lower(?) AND lower(answer)=lower(?)",
                             (q["region"], q["question"], q["answer"])).fetchone()[0]
        if voters < config.AUTO_LEARN_MIN_VOTES and not trusted:
            return
        same = con.execute("SELECT id FROM learned WHERE region=? AND lower(question)=lower(?)",
                           (q["region"], question)).fetchone()
        if same:
            con.execute("UPDATE learned SET votes=COALESCE(votes,0)+1, active=1 WHERE id=?", (same["id"],))
            return
    vec = None
    if llm.embeddings_on():
        try:
            vec = _vec_to_blob(llm.embed([f"{question}\n{q['answer']}"])[0])
        except llm.AIError:
            pass
    with _conn() as con:
        con.execute("INSERT INTO learned (region, question, answer, added_by, added_at, from_question,"
                    " source, active, votes, embedding) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (q["region"], question, q["answer"], "OpsMate (staff 👍)", _now(), q["id"],
                     "auto", 1, 1, vec))


def sources_from_documents(q):
    srcs = json.loads(q["sources"] or "[]")
    return bool(srcs) and srcs[0].get("kind") == "doc"


def review_queue(region=None):
    q = "SELECT * FROM questions WHERE status='needs_review'"
    args = []
    if region:
        q += " AND region=?"
        args.append(region)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def teach(region, question, answer, admin, from_question=None):
    """Store a VERIFIED answer (written by the process team). Used first from now on."""
    vec = None
    if llm.embeddings_on():
        try:
            vec = _vec_to_blob(llm.embed([f"{question}\n{answer}"])[0])
        except llm.AIError:
            pass
    with _conn() as con:
        con.execute("INSERT INTO learned (region, question, answer, added_by, added_at, from_question,"
                    " source, active, votes, embedding) VALUES (?,?,?,?,?,?,?,?,?,?)",
                    (region, question, answer, admin, _now(), from_question, "verified", 1, 0, vec))
        if from_question:
            con.execute("UPDATE questions SET status='resolved', gap_status='Fixed' WHERE id=?",
                        (from_question,))


def dismiss(qid):
    with _conn() as con:
        con.execute("UPDATE questions SET status='dismissed', gap_status='Closed' WHERE id=?", (qid,))


def list_learned(region=None, source=None):
    q, args, where = "SELECT * FROM learned", [], []
    if region:
        where.append("region=?"); args.append(region)
    if source:
        where.append("COALESCE(source,'verified')=?"); args.append(source)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def set_learned(lid, active=None, delete=False, make_verified=False, admin=""):
    with _conn() as con:
        if delete:
            con.execute("DELETE FROM learned WHERE id=?", (lid,))
        if active is not None:
            con.execute("UPDATE learned SET active=? WHERE id=?", (1 if active else 0, lid))
        if make_verified:
            con.execute("UPDATE learned SET source='verified', added_by=? WHERE id=?", (admin, lid))


# ===========================================================================
# 7. KNOWLEDGE GAPS  (questions OpsMate could not answer well)
# ===========================================================================

GAP_STATUSES = ["Open", "SOP update needed", "In progress", "Possibly answered", "Fixed", "Closed"]


def _gap_rows(region=None, since=None):
    q = ("SELECT * FROM questions WHERE (result IN ('not_found','conflict') OR feedback=-1 "
         "OR (result IS NULL AND status='needs_review'))")
    args = []
    if region:
        q += " AND region=?"; args.append(region)
    if since:
        q += " AND asked_at>=?"; args.append(since)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def knowledge_gaps(region=None, since=None, include_closed=False):
    """Group similar unanswered / conflicting / 👎 questions so the process team
    can see what is missing and fix the SOPs. Most-asked first."""
    rows = _gap_rows(region, since)
    if not include_closed:
        rows = [r for r in rows if (r.get("gap_status") or "Open") not in ("Fixed", "Closed")]
    if not rows:
        return []
    texts = [re.sub(r"\s*\[screenshot\]$", "", r["question"]) for r in rows]
    vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(texts)
    m = vec.transform(texts)
    sims = cosine_similarity(m)
    groups, used = [], set()
    for i in range(len(rows)):
        if i in used:
            continue
        members = [j for j in range(len(rows)) if j not in used and (j == i or sims[i, j] >= 0.5)]
        used.update(members)
        g = [rows[j] for j in members]
        kinds = []
        for r in g:
            kinds.append("Conflict" if r.get("result") == "conflict" else
                         "Not helpful" if r.get("feedback") == -1 else "No answer")
        groups.append({
            "question": texts[i], "count": len(g), "ids": [r["id"] for r in g],
            "regions": sorted({r["region"] for r in g}), "topic": g[0].get("topic") or "Other",
            "type": max(set(kinds), key=kinds.count), "first": min(r["asked_at"] for r in g),
            "last": max(r["asked_at"] for r in g), "users": sorted({r["asked_by"] for r in g}),
            "status": g[0].get("gap_status") or "Open", "owner": g[0].get("gap_owner") or "",
            "note": g[0].get("gap_note") or "",
            "examples": list(dict.fromkeys(re.sub(r"\s*\[screenshot\]$", "", r["question"]) for r in g))[:5]})
    groups.sort(key=lambda x: x["last"], reverse=True)       # newest first ...
    groups.sort(key=lambda x: x["count"], reverse=True)      # ... within most-asked first
    return groups


def set_gap(ids, status=None, owner=None, note=None):
    with _conn() as con:
        for i in ids:
            if status is not None:
                con.execute("UPDATE questions SET gap_status=? WHERE id=?", (status, i))
                if status in ("Fixed", "Closed"):
                    con.execute("UPDATE questions SET status='resolved' WHERE id=? AND status='needs_review'", (i,))
            if owner is not None:
                con.execute("UPDATE questions SET gap_owner=? WHERE id=?", (owner, i))
            if note is not None:
                con.execute("UPDATE questions SET gap_note=? WHERE id=?", (note, i))


def recheck_gaps(region, doc_title=""):
    """After a new document is uploaded, test the open gaps again (offline, no AI cost).
    Gaps that now find a good match are marked 'Possibly answered'."""
    try:
        groups = knowledge_gaps(None if region == config.GLOBAL_REGION else region)
    except Exception:
        return 0
    n = 0
    for g in groups:
        if g["status"] not in ("Open", "SOP update needed", "In progress") or g["type"] == "Not helpful":
            continue
        for r in g["regions"]:
            hits = search(r, [brain.search_text(g["question"])])
            if hits and hits[0]["kind"] == "doc" and hits[0]["title"] == doc_title:
                st = brain.core_text(g["question"])
                conf = brain.offline_confidence(hits[0]["lex"], brain.coverage(st, hits[0]["search"]),
                                                None, len(brain.content_words(st)))
                if conf != "low":
                    set_gap(g["ids"], status="Possibly answered",
                            note=f"New document '{doc_title}' may answer this - ask again to confirm.")
                    n += 1
                    break
    return n


# ===========================================================================
# 8. HISTORY, ANALYTICS AND EXCEL EXPORT
# ===========================================================================

def history(region=None, user=None, limit=50):
    q, args, where = "SELECT * FROM questions", [], []
    if region:
        where.append("region=?"); args.append(region)
    if user:
        where.append("asked_by=?"); args.append(user)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + f" ORDER BY id DESC LIMIT {int(limit)}", args)]


def questions_df(region=None, days=None, topic=None):
    """All questions as a table for the dashboard and the Excel export."""
    import pandas as pd
    q, args, where = "SELECT * FROM questions", [], []
    if region and region != "All":
        where.append("region=?"); args.append(region)
    if days:
        where.append("asked_at>=?")
        args.append((datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d %H:%M"))
    if topic and topic != "All":
        where.append("COALESCE(topic,'Other')=?"); args.append(topic)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        df = pd.read_sql_query(q + " ORDER BY id DESC", con, params=args)
    if df.empty:
        return df
    df["topic"] = df["topic"].fillna("Other")
    df["result"] = df.apply(lambda r: r["result"] if r["result"] else
                            ("answered" if r["status"] in ("answered", "resolved") else "not_found"), axis=1)
    df["confidence"] = df["confidence"].fillna("")
    df["date"] = df["asked_at"].str.slice(0, 10)
    df["helpful"] = df["feedback"].map({1: "Yes", -1: "No"}).fillna("")
    df["first_source"] = df["sources"].apply(
        lambda s: (json.loads(s)[0]["title"] if s and s != "[]" else ""))
    return df


def export_excel(df, gaps):
    """Excel workbook: Summary, All Questions, Unanswered & Conflicts, Knowledge Gaps, Feedback."""
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    buf = io.BytesIO()
    cols = {"asked_at": "Asked at", "region": "Region", "asked_by": "Asked by", "topic": "Topic",
            "question": "Question", "result": "Result", "confidence": "Confidence",
            "answer": "Answer given", "first_source": "Main source", "helpful": "Helpful?",
            "comment": "Staff comment", "status": "Workflow status", "gap_status": "Gap status"}
    qdf = df[list(cols)].rename(columns=cols) if not df.empty else pd.DataFrame(columns=list(cols.values()))
    total = len(df)
    ans = int((df["result"] == "answered").sum()) if total else 0
    rated = int(df["feedback"].notna().sum()) if total else 0
    summary = pd.DataFrame([
        ["Total questions", total],
        ["Answered", ans], ["Answered %", f"{100 * ans / total:.0f}%" if total else "-"],
        ["No answer (sent to supervisor)", int((df["result"] == "not_found").sum()) if total else 0],
        ["Conflicting documents", int((df["result"] == "conflict").sum()) if total else 0],
        ["Rated helpful", int((df["feedback"] == 1).sum()) if total else 0],
        ["Rated not helpful", int((df["feedback"] == -1).sum()) if total else 0],
        ["Helpful % (of rated)", f"{100 * (df['feedback'] == 1).sum() / rated:.0f}%" if rated else "-"],
        ["Open knowledge gaps", len(gaps)],
        ["Exported at", _now()]], columns=["Measure", "Value"])
    gdf = pd.DataFrame([{"Times asked": g["count"], "Question": g["question"], "Type": g["type"],
                         "Topic": g["topic"], "Regions": ", ".join(g["regions"]),
                         "First asked": g["first"], "Last asked": g["last"], "Status": g["status"],
                         "Owner": g["owner"], "Note": g["note"],
                         "Similar wordings": " | ".join(g["examples"])} for g in gaps])
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        summary.to_excel(xw, sheet_name="Summary", index=False)
        qdf.to_excel(xw, sheet_name="All Questions", index=False)
        qdf[qdf["Result"].isin(["not_found", "conflict"])].to_excel(xw, sheet_name="Unanswered & Conflicts", index=False)
        (gdf if not gdf.empty else pd.DataFrame(columns=["Times asked", "Question"])).to_excel(
            xw, sheet_name="Knowledge Gaps", index=False)
        qdf[qdf["Helpful?"] != ""].to_excel(xw, sheet_name="Feedback", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F5FBF")
                c.alignment = Alignment(vertical="center")
            ws.freeze_panes = "A2"
            if ws.max_row > 1:
                ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                longest = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(60, max(12, longest + 2))
                for c in col[1:]:
                    c.alignment = Alignment(wrap_text=longest > 60, vertical="top")
    return buf.getvalue()


# ===========================================================================
# 9. ANNOUNCEMENTS  ("Important Updates" on the home screen)
# ===========================================================================

def add_announcement(title, body, level, region, author):
    with _conn() as con:
        con.execute("INSERT INTO announcements (title, body, level, region, author, created_at, active)"
                    " VALUES (?,?,?,?,?,?,1)", (title, body, level, region, author, _now()))


def announcements(region=None, active_only=True):
    q, args = "SELECT * FROM announcements", []
    where = []
    if active_only:
        where.append("active=1")
    if region:
        where.append("region IN (?, 'All')"); args.append(region)
    if where:
        q += " WHERE " + " AND ".join(where)
    with _conn() as con:
        return [dict(r) for r in con.execute(q + " ORDER BY id DESC", args)]


def set_announcement(aid, active):
    with _conn() as con:
        con.execute("UPDATE announcements SET active=? WHERE id=?", (1 if active else 0, aid))
